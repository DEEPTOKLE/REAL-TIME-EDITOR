"""
Tests for the .docx import/export bridge.
"""

import asyncio
import io
import random
import shutil
import tempfile
import unittest

import aiohttp
from aiohttp import web
from docx import Document
from docx.shared import Pt

from backend.docx_bridge import docx_to_synapse, synapse_to_docx
from backend.server import CollaborativeServer


def _make_docx(paragraphs):
    """Helper: build a .docx in memory from a list of paragraph dicts.

    Each paragraph dict: {"style": "Normal"/"Heading 1"/..., "runs": [{"text":..., "bold":..., "italic":..., "underline":...}, ...]}
    """
    doc = Document()
    for p in paragraphs:
        style = p.get("style", "Normal")
        try:
            para = doc.add_paragraph(style=style)
        except Exception:
            para = doc.add_paragraph()
        for run in p.get("runs", []):
            r = para.add_run(run.get("text", ""))
            r.bold = bool(run.get("bold"))
            r.italic = bool(run.get("italic"))
            r.underline = bool(run.get("underline"))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestDocxBridge(unittest.TestCase):
    def test_plain_text_roundtrip(self):
        docx_bytes = _make_docx([
            {"style": "Normal", "runs": [{"text": "Hello world"}]},
            {"style": "Normal", "runs": [{"text": "Second paragraph"}]},
        ])
        text, intervals = docx_to_synapse(docx_bytes)
        self.assertEqual(text, "Hello world\nSecond paragraph")
        # Round-trip through export
        out = synapse_to_docx(text, intervals)
        doc2 = Document(io.BytesIO(out))
        self.assertEqual(len(doc2.paragraphs), 2)
        self.assertEqual(doc2.paragraphs[0].text, "Hello world")
        self.assertEqual(doc2.paragraphs[1].text, "Second paragraph")

    def test_bold_italic_underline_preserved(self):
        docx_bytes = _make_docx([
            {
                "style": "Normal",
                "runs": [
                    {"text": "plain "},
                    {"text": "bold", "bold": True},
                    {"text": " plain "},
                    {"text": "italic", "italic": True},
                    {"text": " "},
                    {"text": "under", "underline": True},
                ],
            }
        ])
        text, intervals = docx_to_synapse(docx_bytes)
        self.assertEqual(text, "plain bold plain italic under")
        # Build a map of position -> attrs
        attrs_by_pos = {}
        for s, e, a in intervals:
            for i in range(s, e):
                attrs_by_pos[i] = a
        self.assertEqual(attrs_by_pos[6], {"bold": True})
        self.assertEqual(attrs_by_pos[17], {"italic": True})
        self.assertEqual(attrs_by_pos[24], {"underline": True})
        # Positions with no formatting should not appear in attrs_by_pos
        self.assertNotIn(0, attrs_by_pos)

    def test_heading_styles_mapped(self):
        docx_bytes = _make_docx([
            {"style": "Heading 1", "runs": [{"text": "Title"}]},
            {"style": "Heading 2", "runs": [{"text": "Subtitle"}]},
            {"style": "Heading 3", "runs": [{"text": "Section"}]},
        ])
        text, intervals = docx_to_synapse(docx_bytes)
        self.assertEqual(text, "Title\nSubtitle\nSection")
        # First chars of each line should carry the header attribute
        attrs_by_pos = {}
        for s, e, a in intervals:
            for i in range(s, e):
                attrs_by_pos[i] = a
        self.assertEqual(attrs_by_pos[0], {"header": 1})
        self.assertEqual(attrs_by_pos[6], {"header": 2})
        self.assertEqual(attrs_by_pos[15], {"header": 3})

    def test_export_produces_valid_docx(self):
        text = "Hello world\nSecond line"
        intervals = [[0, 5, {"bold": True}], [12, 17, {"italic": True}]]
        out = synapse_to_docx(text, intervals)
        doc = Document(io.BytesIO(out))
        self.assertEqual(len(doc.paragraphs), 2)
        self.assertEqual(doc.paragraphs[0].text, "Hello world")
        run_bold = doc.paragraphs[0].runs[0]
        self.assertTrue(run_bold.bold)
        run_italic = doc.paragraphs[1].runs[0]
        self.assertTrue(run_italic.italic)

    def test_invalid_file_raises_value_error(self):
        with self.assertRaises(ValueError):
            docx_to_synapse(b"this is not a docx file")

    def test_blockquote_mapped(self):
        docx_bytes = _make_docx([
            {"style": "Quote", "runs": [{"text": "quoted text"}]},
        ])
        text, intervals = docx_to_synapse(docx_bytes)
        self.assertTrue(any("blockquote" in a for _, _, a in intervals))

    def test_synapse_to_docx_roundtrip_preserves_text(self):
        text = "Line one\nLine two\nLine three"
        out = synapse_to_docx(text, [])
        doc = Document(io.BytesIO(out))
        self.assertEqual([p.text for p in doc.paragraphs], ["Line one", "Line two", "Line three"])

    def test_image_extracted_from_docx(self):
        docx_bytes = _make_docx([
            {"style": "Normal", "runs": [{"text": "Before "}]},
            {"style": "Normal", "runs": [{"text": "", "image": True}]},
            {"style": "Normal", "runs": [{"text": " After"}]},
        ])
        # _make_docx doesn't support images; we need a real docx with an image
        doc = Document()
        para = doc.add_paragraph("Before ")
        run = para.add_run()
        run.add_picture(io.BytesIO(b'GIF89a\x01\x00\x01\x00\x80\x00\x00\x00\x00\x00\x00\x00!\xf9\x04\x01\x00\x00\x00\x00,\x00\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02D\x01\x00;'), width=Pt(100))
        para.add_run(" After")
        buf = io.BytesIO()
        doc.save(buf)
        text, intervals = docx_to_synapse(buf.getvalue())
        self.assertIn("\uFFFC", text)
        img_intervals = [iv for iv in intervals if isinstance(iv[2], dict) and iv[2].get("image")]
        self.assertEqual(len(img_intervals), 1)
        self.assertEqual(img_intervals[0][0], img_intervals[0][1] - 1)

    def test_image_roundtrip(self):
        doc = Document()
        para = doc.add_paragraph("Text ")
        run = para.add_run()
        run.add_picture(io.BytesIO(b'GIF89a\x01\x00\x01\x00\x80\x00\x00\x00\x00\x00\x00\x00!\xf9\x04\x01\x00\x00\x00\x00,\x00\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02D\x01\x00;'), width=Pt(100))
        para.add_run(" more")
        buf = io.BytesIO()
        doc.save(buf)

        text, intervals = docx_to_synapse(buf.getvalue())
        out = synapse_to_docx(text, intervals)
        doc2 = Document(io.BytesIO(out))
        self.assertEqual(len(doc2.inline_shapes), 1)

    def test_image_store_save_and_load(self):
        from backend.image_store import save_image_file, load_image_file
        data = b'GIF89a\x01\x00\x01\x00\x80\x00\x00\x00\x00\x00\x00\x00!\xf9\x04\x01\x00\x00\x00\x00,\x00\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02D\x01\x00;'
        img_id = save_image_file(data, "image/gif")
        loaded, ctype = load_image_file(img_id)
        self.assertEqual(data, loaded)
        self.assertEqual(ctype, "image/gif")

    def test_oversized_image_rejected(self):
        from backend.image_store import save_image_file
        with self.assertRaises(ValueError):
            save_image_file(b"x" * (10 * 1024 * 1024 + 1), "image/png")



import asyncio
import aiohttp
from aiohttp import web
from backend.server import CollaborativeServer


class TestDocxBridgeServerIntegration(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.mkdtemp()
        self.port = 8796
        self.host = "127.0.0.1"
        self.server_obj = CollaborativeServer(host=self.host, port=self.port, data_dir=self.tmp)
        self.runner = web.AppRunner(self.server_obj.app)
        await self.runner.setup()
        self.site = web.TCPSite(self.runner, self.host, self.port)
        await self.site.start()
        self.http_url = f"http://{self.host}:{self.port}"

    async def asyncTearDown(self):
        await self.runner.cleanup()
        shutil.rmtree(self.tmp, ignore_errors=True)

    async def test_doc_extension_rejected(self):
        payload = b"PK\x03\x04fake-docx-bytes"
        data = aiohttp.FormData()
        data.add_field(
            "file",
            payload,
            filename="uploaded.doc",
            content_type="application/octet-stream",
        )
        async with aiohttp.ClientSession() as session:
            async with session.post(f"{self.http_url}/api/documents/import", data=data) as resp:
                self.assertEqual(resp.status, 400)
                body = await resp.json()
                self.assertIn("Only .docx files are supported", body.get("error", ""))


if __name__ == "__main__":
    unittest.main()
