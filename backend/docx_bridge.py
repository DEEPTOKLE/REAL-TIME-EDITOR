"""
Word (.docx) ↔ Synapse document conversion bridge.

All Word read/write operations are isolated here so that the rest of the
codebase (OT, Rope, WebSocket, formatting store) stays untouched.
"""

import io
from typing import List, Tuple, Any

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from .image_store import save_image_file, load_image_file, ALLOWED_EXTENSIONS, CONTENT_TYPE_ALIASES

ALLOWED_CONTENT_TYPES = set(ALLOWED_EXTENSIONS) | set(CONTENT_TYPE_ALIASES)


def _get_run_formatting(run_elem):
    """Extract bold, italic, underline from a w:r XML element."""
    rPr = run_elem.find(qn("w:rPr"))
    if rPr is None:
        return {}
    attrs = {}
    for attr_name, tag in [("bold", "w:b"), ("italic", "w:i"), ("underline", "w:u")]:
        elem = rPr.find(qn(tag))
        if elem is not None:
            val = elem.get(qn("w:val"))
            if val is None or val.lower() not in ("0", "false", "none"):
                attrs[attr_name] = True
    return attrs


def _get_image_info(para, rel_id):
    """Return (content_type, width_px, height_px) for an inline image."""
    try:
        part = para.part
        image_part = part.related_parts.get(rel_id)
        if image_part is None:
            return None, 300, 200

        content_type = image_part.content_type
        if content_type not in ALLOWED_CONTENT_TYPES:
            return None, 300, 200

        width_px = 300
        height_px = 200

        for child in para._p:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "r":
                drawing = child.find(qn("w:drawing"))
                if drawing is not None:
                    blip = drawing.find(".//" + qn("a:blip"))
                    if blip is not None and blip.get(qn("r:embed")) == rel_id:
                        extent = drawing.find(".//" + qn("wp:extent"))
                        if extent is not None:
                            cx = int(extent.get("cx", 0))
                            cy = int(extent.get("cy", 0))
                            if cx > 0 and cy > 0:
                                width_px = cx / 914400 * 96
                                height_px = cy / 914400 * 96
                        break

        return content_type, int(width_px), int(height_px)
    except Exception:
        return None, 300, 200


def docx_to_synapse(file_bytes: bytes) -> Tuple[str, List[List[Any]]]:
    """Convert a .docx file into Synapse plain text + formatting intervals.

    Args:
        file_bytes: Raw bytes of a .docx file.

    Returns:
        (text, formatting_intervals) where formatting_intervals is a list of
        [start, end, {attr: val}] intervals (only non-empty attribute dicts).

    Raises:
        ValueError: If the file cannot be parsed as a valid .docx document.
    """
    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise ValueError(f"Failed to parse .docx file: {exc}") from exc

    text_parts: List[str] = []
    intervals: List[List[Any]] = []
    pos = 0

    # Heading / block style → attribute mapping at line start
    HEADING_STYLES = {"Heading 1": 1, "Heading 2": 2, "Heading 3": 3}
    BLOCKQUOTE_STYLES = {"Quote", "Block Text"}

    for para_idx, para in enumerate(doc.paragraphs):
        line_start = pos
        line_len = 0

        # Determine block attribute for this paragraph
        block_attr: dict = {}
        style_name = para.style.name if para.style else ""
        if style_name in HEADING_STYLES:
            block_attr["header"] = HEADING_STYLES[style_name]
        elif style_name in BLOCKQUOTE_STYLES:
            block_attr["blockquote"] = True

        # Process XML elements in document order
        for child in para._p:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

            if tag == "r":
                run_attrs = _get_run_formatting(child)
                t_elem = child.find(qn("w:t"))
                if t_elem is not None and t_elem.text:
                    run_text = t_elem.text
                    start = pos
                    end = pos + len(run_text)
                    merged_attrs = {**block_attr, **run_attrs}
                    if merged_attrs:
                        intervals.append([start, end, merged_attrs])
                    text_parts.append(run_text)
                    pos += len(run_text)
                    line_len += len(run_text)

                drawing = child.find(qn("w:drawing"))
                if drawing is not None:
                    blip = drawing.find(".//" + qn("a:blip"))
                    if blip is not None:
                        embed = blip.get(qn("r:embed"))
                        if embed:
                            content_type, width_px, height_px = _get_image_info(para, embed)
                            if content_type:
                                try:
                                    image_part = para.part.related_parts.get(embed)
                                    if image_part:
                                        img_id = save_image_file(image_part.blob, content_type)
                                        img_attrs = {
                                            "image": img_id,
                                            "width": width_px,
                                            "height": height_px,
                                        }
                                        if block_attr:
                                            img_attrs = {**block_attr, **img_attrs}
                                        intervals.append([pos, pos + 1, img_attrs])
                                        text_parts.append("\uFFFC")
                                        pos += 1
                                        line_len += 1
                                except Exception:
                                    pass

            elif tag == "hyperlink":
                for r in child.findall(qn("w:r")):
                    run_attrs = _get_run_formatting(r)
                    t_elem = r.find(qn("w:t"))
                    if t_elem is not None and t_elem.text:
                        run_text = t_elem.text
                        start = pos
                        end = pos + len(run_text)
                        merged_attrs = {**block_attr, **run_attrs}
                        if merged_attrs:
                            intervals.append([start, end, merged_attrs])
                        text_parts.append(run_text)
                        pos += len(run_text)
                        line_len += len(run_text)

                    drawing = r.find(qn("w:drawing"))
                    if drawing is not None:
                        blip = drawing.find(".//" + qn("a:blip"))
                        if blip is not None:
                            embed = blip.get(qn("r:embed"))
                            if embed:
                                content_type, width_px, height_px = _get_image_info(para, embed)
                                if content_type:
                                    try:
                                        image_part = para.part.related_parts.get(embed)
                                        if image_part:
                                            img_id = save_image_file(image_part.blob, content_type)
                                            img_attrs = {
                                                "image": img_id,
                                                "width": width_px,
                                                "height": height_px,
                                            }
                                            if block_attr:
                                                img_attrs = {**block_attr, **img_attrs}
                                            intervals.append([pos, pos + 1, img_attrs])
                                            text_parts.append("\uFFFC")
                                            pos += 1
                                            line_len += 1
                                    except Exception:
                                        pass

        # If the paragraph had a block attribute but no runs emitted it
        # (e.g. empty heading), still record the block interval.
        if block_attr and line_len == 0:
            intervals.append([line_start, line_start, block_attr])

        # Add newline between paragraphs (not after the very last one)
        if para_idx < len(doc.paragraphs) - 1:
            text_parts.append("\n")
            pos += 1

    text = "".join(text_parts)
    return text, intervals


def synapse_to_docx(text: str, formatting_intervals: List[List[Any]]) -> bytes:
    """Convert Synapse text + formatting intervals into a .docx file.

    Args:
        text: Plain text document content.
        formatting_intervals: List of [start, end, {attr: val}] intervals.

    Returns:
        Raw bytes of the generated .docx file.

    Raises:
        ValueError: If the document cannot be written.
    """
    try:
        doc = Document()
    except Exception as exc:
        raise ValueError(f"Failed to create Word document: {exc}") from exc

    HEADING_STYLES = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}

    # Build per-character attribute lookup
    attr_by_index = {}
    for start, end, attrs in formatting_intervals:
        if not attrs:
            continue
        for i in range(start, end):
            if i < len(text):
                attr_by_index.setdefault(i, {}).update(attrs)

    lines = text.split("\n")
    pos = 0

    for line_idx, line in enumerate(lines):
        line_len = len(line)

        # Determine block style from the first character's attributes
        block_style = None
        if line_len > 0:
            first_attrs = attr_by_index.get(pos, {})
            header_level = first_attrs.get("header")
            if header_level and header_level in HEADING_STYLES:
                block_style = HEADING_STYLES[header_level]
            elif first_attrs.get("blockquote"):
                block_style = "Quote"

        try:
            if block_style:
                para = doc.add_paragraph(style=block_style)
            else:
                para = doc.add_paragraph()
        except Exception:
            para = doc.add_paragraph()

        # Process character by character, handling images
        i = 0
        while i < line_len:
            if line[i] == "\uFFFC":
                img_attrs = attr_by_index.get(pos + i, {})
                img_id = img_attrs.get("image")
                if img_id:
                    try:
                        img_bytes, content_type = load_image_file(img_id)
                        width_px = img_attrs.get("width", 300) or 300
                        height_px = img_attrs.get("height", 200) or 200
                        width_pt = width_px * 0.75
                        height_pt = height_px * 0.75
                        run = para.add_run()
                        run.add_picture(io.BytesIO(img_bytes), width=Pt(width_pt), height=Pt(height_pt))
                    except Exception:
                        run = para.add_run("[Image unavailable]")
                else:
                    run = para.add_run("\uFFFC")
                i += 1
                continue

            # Normal text - group consecutive chars with same attrs
            attrs = attr_by_index.get(pos + i, {})
            j = i + 1
            while j < line_len and line[j] != "\uFFFC" and attr_by_index.get(pos + j, {}) == attrs:
                j += 1

            chunk = line[i:j]
            run = para.add_run(chunk)
            run.bold = bool(attrs.get("bold"))
            run.italic = bool(attrs.get("italic"))
            run.underline = bool(attrs.get("underline"))

            i = j

        pos += line_len + 1  # +1 for the \n we split on

    buffer = io.BytesIO()
    try:
        doc.save(buffer)
    except Exception as exc:
        raise ValueError(f"Failed to save Word document: {exc}") from exc

    return buffer.getvalue()
